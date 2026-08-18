from pathlib import Path
import json
from openpyxl import load_workbook

root = Path("/home/ubuntu/appointment/data/who")
sampled = {}
for path in sorted(root.glob("*.xlsx")):
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    key_to_index = {name: index for index, name in enumerate(headers)}
    target_months = {24, 30, 36, 42, 48, 54, 60}
    sampled[path.stem] = [
        {"month": row[key_to_index["Month"]], "p3": row[key_to_index["P3"]], "p50": row[key_to_index["P50"]], "p97": row[key_to_index["P97"]]}
        for row in rows[1:]
        if row[key_to_index["Month"]] in target_months
    ]
print(json.dumps(sampled, separators=(",", ":")))
